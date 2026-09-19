from __future__ import annotations

from pathlib import Path

from harbormaster.reader.parsers.base import BaseParser


class XlsxParser(BaseParser):
    name = "xlsx"

    def parse(self, path: Path) -> str:
        from openpyxl import load_workbook

        wb = load_workbook(path, read_only=True, data_only=True)
        lines: list[str] = []
        for sheet in wb.worksheets:
            for row in sheet.iter_rows(values_only=True):
                cells = [str(c).strip() for c in row if c is not None and str(c).strip()]
                if not cells:
                    continue
                if len(cells) == 1:
                    lines.append(cells[0])
                else:
                    lines.append(f"{cells[0]}: {cells[1]}")
        wb.close()
        return "\n".join(lines)
