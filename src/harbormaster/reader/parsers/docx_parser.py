from __future__ import annotations

from pathlib import Path

from harbormaster.reader.parsers.base import BaseParser


class DocxParser(BaseParser):
    name = "docx"

    def parse(self, path: Path) -> str:
        from docx import Document

        doc = Document(path)
        lines = [p.text for p in doc.paragraphs if p.text]
        for table in doc.tables:
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells if c.text and c.text.strip()]
                if not cells:
                    continue
                if len(cells) == 1:
                    lines.append(cells[0])
                else:
                    lines.append(f"{cells[0]}: {cells[1]}")
        return "\n".join(lines)
