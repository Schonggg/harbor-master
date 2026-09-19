"""Field x format x label matrix. Every cell is a real extract + L5 exact compare.

Scoring is asymmetric: a false alarm zeros the record. A miss can still be saved.
Never add fuzzy / similarity / a third compare state to close a cell.
"""

from __future__ import annotations

from pathlib import Path
from typing import Any, Iterator

from harbormaster.config import get_field_aliases
from harbormaster.models import COMPARE_FIELDS
from harbormaster.official.compare import values_match
from harbormaster.reader.extractor import FieldExtractor

FORMATS = ("txt", "pdf", "docx", "xlsx")

GOLD = {
    "shipper": "ACME TRADING CO., LTD.",
    "consignee": "BETA IMPORTS PTE LTD",
    "notify_party": "SAME AS CONSIGNEE",
    "port_of_loading": "Shanghai (CNSHA)",
    "port_of_discharge": "Los Angeles (USLAX)",
    "container_count": "THREE (3) x 40HC",
    "gross_weight_kg": "1,000 KGS",
}

# Same facts, different writing. L5 must still MATCH — these are the false-alarm traps.
EQUIVALENT = {
    "shipper": "ACME TRADING CO LTD",
    "consignee": "Beta Imports Pte Ltd",
    "notify_party": "same as consignee",
    "port_of_loading": "shanghai",
    "port_of_discharge": "los angeles",
    "container_count": "3",
    "gross_weight_kg": "1000",
}

# Real discrepancies. L5 must still fire.
NEAR_MISS = {
    "shipper": "ACME TRADING INC",
    "consignee": "BETA EXPORTS PTE LTD",
    "notify_party": "GAMMA LOGISTICS",
    "port_of_loading": "Ningbo",
    "port_of_discharge": "Long Beach",
    "container_count": "4",
    "gross_weight_kg": "14800",
}

TALKING_POINT = (
    "The score sheet punishes false alarms harder than misses. We do not sample a few "
    "txt cases. Every field x txt/pdf/docx/xlsx x known label is a cell, and every "
    "normalization is checked against equivalent pairs before it is allowed to catch a mismatch."
)


def labels_for(field: str) -> list[str]:
    names = [field.replace("_", " ")]
    for alias in get_field_aliases().get(field, []):
        if alias and alias not in names:
            names.append(alias)
    return names


def cells() -> list[dict[str, str]]:
    out: list[dict[str, str]] = []
    for field in COMPARE_FIELDS:
        for fmt in FORMATS:
            for label in labels_for(field):
                out.append({"field": field, "fmt": fmt, "label": label})
    return out


def iter_cells() -> Iterator[dict[str, str]]:
    yield from cells()


def _write_txt(path: Path, label: str, value: str) -> None:
    path.write_text(f"{label}: {value}\n", encoding="utf-8")


def _write_pdf(path: Path, label: str, value: str) -> None:
    import pymupdf

    doc = pymupdf.open()
    page = doc.new_page()
    rect = pymupdf.Rect(48, 48, 560, 200)
    page.insert_textbox(rect, f"{label}: {value}", fontsize=11, fontname="helv")
    doc.save(path)
    doc.close()


def _write_docx(path: Path, label: str, value: str) -> None:
    from docx import Document

    doc = Document()
    doc.add_paragraph(f"{label}: {value}")
    doc.save(path)


def _write_xlsx(path: Path, label: str, value: str) -> None:
    from openpyxl import Workbook

    wb = Workbook()
    sheet = wb.active
    sheet["A1"] = label
    sheet["B1"] = value
    wb.save(path)


_WRITERS = {
    "txt": _write_txt,
    "pdf": _write_pdf,
    "docx": _write_docx,
    "xlsx": _write_xlsx,
}


def write_cell(folder: Path, field: str, fmt: str, label: str, value: str) -> Path:
    safe = "".join(ch if ch.isalnum() else "_" for ch in f"{field}_{fmt}_{label}")[:80]
    path = folder / f"{safe}.{fmt}"
    _WRITERS[fmt](path, label, value)
    return path


def run_cell(folder: Path, field: str, fmt: str, label: str) -> dict[str, Any]:
    path = write_cell(folder, field, fmt, label, GOLD[field])
    extractor = FieldExtractor(degrade=True, rules_only=True)
    doc = extractor.extract_path(path)
    raw = (doc.fields.get(field).raw_value if field in doc.fields else "") or ""
    ok = bool(raw) and values_match(field, GOLD[field], raw)
    return {
        "field": field,
        "fmt": fmt,
        "label": label,
        "ok": ok,
        "extracted": raw,
        "parser": doc.parser_used,
    }


def run_matrix(folder: Path) -> dict[str, Any]:
    extractor = FieldExtractor(degrade=True, rules_only=True)
    rows = []
    for c in iter_cells():
        path = write_cell(folder, c["field"], c["fmt"], c["label"], GOLD[c["field"]])
        doc = extractor.extract_path(path)
        raw = (doc.fields.get(c["field"]).raw_value if c["field"] in doc.fields else "") or ""
        ok = bool(raw) and values_match(c["field"], GOLD[c["field"]], raw)
        rows.append(
            {
                "field": c["field"],
                "fmt": c["fmt"],
                "label": c["label"],
                "ok": ok,
                "extracted": raw,
                "parser": doc.parser_used,
            }
        )
    failed = [r for r in rows if not r["ok"]]
    by_fmt: dict[str, dict[str, int]] = {fmt: {"ok": 0, "fail": 0} for fmt in FORMATS}
    by_field: dict[str, dict[str, int]] = {f: {"ok": 0, "fail": 0} for f in COMPARE_FIELDS}
    grid: dict[str, dict[str, dict[str, int]]] = {
        f: {fmt: {"ok": 0, "fail": 0, "n": 0} for fmt in FORMATS} for f in COMPARE_FIELDS
    }
    for row in rows:
        bucket = "ok" if row["ok"] else "fail"
        by_fmt[row["fmt"]][bucket] += 1
        by_field[row["field"]][bucket] += 1
        cell = grid[row["field"]][row["fmt"]]
        cell["n"] += 1
        cell[bucket] += 1
    return {
        "cells": len(rows),
        "passed": len(rows) - len(failed),
        "failed": failed,
        "by_fmt": by_fmt,
        "by_field": by_field,
        "grid": grid,
        "formats": list(FORMATS),
        "fields": list(COMPARE_FIELDS),
    }


def false_alarm_pairs() -> list[dict[str, Any]]:
    rows = []
    for field in COMPARE_FIELDS:
        match_ok = values_match(field, GOLD[field], EQUIVALENT[field])
        miss_ok = not values_match(field, GOLD[field], NEAR_MISS[field])
        rows.append(
            {
                "field": field,
                "equivalent_ok": match_ok,
                "near_miss_ok": miss_ok,
                "false_alarm": not match_ok,
            }
        )
    return rows


def would_false_alarm(field: str, left: str, right: str) -> bool:
    """True if two writings of the same fact would be scored as a defect."""
    return not values_match(field, left, right)
