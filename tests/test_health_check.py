"""Each of the four NEEDS_REVIEW triggers."""

from __future__ import annotations

from harbormaster.models import EmailMessage, ExtractedDocument, FieldValue, ReviewReason
from harbormaster.reader.health_check import classify_doc_kind, health_check


def _fv(name: str, value: str) -> FieldValue:
    return FieldValue(name=name, raw_value=value, confidence=0.95)


def _full_fields() -> dict[str, FieldValue]:
    return {
        "shipper": _fv("shipper", "ACME"),
        "consignee": _fv("consignee", "BETA"),
        "notify_party": _fv("notify_party", "SAME AS CONSIGNEE"),
        "port_of_loading": _fv("port_of_loading", "Shanghai"),
        "port_of_discharge": _fv("port_of_discharge", "Los Angeles"),
        "container_count": _fv("container_count", "3"),
        "gross_weight_kg": _fv("gross_weight_kg", "1000 KGS"),
    }


def test_missing_attachment():
    email = EmailMessage(email_id="e1", subject="please check BL draft vs SI")
    result = health_check(email, [])
    assert result.ok is False
    assert result.reason == ReviewReason.MISSING_ATTACHMENT


def test_unreadable():
    email = EmailMessage(email_id="e1", subject="SI vs BL", attachment_paths=["si.pdf", "bl.pdf"])
    si = ExtractedDocument(filename="si.pdf", kind="si", text="", parser_used="vision")
    bl = ExtractedDocument(filename="bl.pdf", kind="bl", text="", parser_used="vision")
    result = health_check(email, [si, bl], si=si, bl=bl)
    assert result.ok is False
    assert result.reason == ReviewReason.UNREADABLE


def test_wrong_doc_type():
    email = EmailMessage(email_id="e1", subject="SI vs BL", attachment_paths=["inv.pdf", "pl.pdf"])
    fields = _full_fields()
    si = ExtractedDocument(
        filename="commercial_invoice.pdf",
        kind="other",
        text="COMMERCIAL INVOICE\n" + "x" * 80,
        fields=fields,
    )
    bl = ExtractedDocument(
        filename="packing_list.pdf",
        kind="other",
        text="PACKING LIST\n" + "x" * 80,
        fields=fields,
    )
    result = health_check(email, [si, bl], si=si, bl=bl)
    assert result.ok is False
    assert result.reason == ReviewReason.WRONG_DOC_TYPE


def test_missing_value():
    email = EmailMessage(email_id="e1", subject="SI vs BL", attachment_paths=["si.txt", "bl.txt"])
    fields = _full_fields()
    si_fields = dict(fields)
    si_fields.pop("shipper")
    si = ExtractedDocument(
        filename="si.txt",
        kind="si",
        text="SHIPPING INSTRUCTION\nConsignee: BETA\n" + "x" * 80,
        fields=si_fields,
    )
    bl = ExtractedDocument(
        filename="bl.txt",
        kind="bl",
        text="BILL OF LADING\nConsignee: BETA\n" + "x" * 80,
        fields=fields,
    )
    result = health_check(email, [si, bl], si=si, bl=bl)
    assert result.ok is False
    assert result.reason == ReviewReason.MISSING_VALUE
    assert "shipper" in result.missing_fields


def test_healthy_pair_passes():
    email = EmailMessage(email_id="e1", subject="SI vs BL", attachment_paths=["si.txt", "bl.txt"])
    fields = _full_fields()
    si = ExtractedDocument(
        filename="si.txt",
        kind="si",
        text="SHIPPING INSTRUCTION\n" + "x" * 80,
        fields=fields,
    )
    bl = ExtractedDocument(
        filename="bl.txt",
        kind="bl",
        text="BILL OF LADING\n" + "x" * 80,
        fields=fields,
    )
    result = health_check(email, [si, bl], si=si, bl=bl)
    assert result.ok is True
    assert result.reason is None


def _invoice_body() -> str:
    return "\n".join(
        [
            "Item  Qty  Line total",
            "Paper  10  $120.00",
            "Board  20  $240.00",
            "Pulp   5   USD 80.50",
            "Tax    1   $12.00",
        ]
    )


def _packing_body() -> str:
    return "\n".join(
        [
            "Pieces by carton",
            "CTN-01 paper reams",
            "CTN-02 paper reams",
            "CTN-03 pulp sheets",
            "CARTON NO. 4 mixed",
        ]
    )


def test_invoice_structure_without_keywords_is_wrong_doc_type():
    inv = ExtractedDocument(filename="charges.pdf", text=_invoice_body(), fields={})
    assert classify_doc_kind(inv) == "other"
    email = EmailMessage(email_id="e1", subject="SI vs BL", attachment_paths=["si.txt", "charges.pdf"])
    fields = _full_fields()
    si = ExtractedDocument(
        filename="si.txt",
        kind="si",
        text="SHIPPING INSTRUCTION\n" + "x" * 80,
        fields=fields,
    )
    inv.kind = classify_doc_kind(inv)
    result = health_check(email, [si, inv], si=si, bl=inv)
    assert result.ok is False
    assert result.reason == ReviewReason.WRONG_DOC_TYPE


def test_packing_structure_without_keywords_is_wrong_doc_type():
    pl = ExtractedDocument(filename="cartons.pdf", text=_packing_body(), fields={})
    assert classify_doc_kind(pl) == "other"
    email = EmailMessage(email_id="e1", subject="SI vs BL", attachment_paths=["si.txt", "cartons.pdf"])
    fields = _full_fields()
    si = ExtractedDocument(
        filename="si.txt",
        kind="si",
        text="SHIPPING INSTRUCTION\n" + "x" * 80,
        fields=fields,
    )
    pl.kind = classify_doc_kind(pl)
    result = health_check(email, [si, pl], si=si, bl=pl)
    assert result.ok is False
    assert result.reason == ReviewReason.WRONG_DOC_TYPE


def test_normal_bl_is_not_structurally_wrong():
    fields = _full_fields()
    bl = ExtractedDocument(
        filename="bl.txt",
        text="BILL OF LADING\nPort of loading Shanghai\n3 x 40HC\n" + "x" * 80,
        fields=fields,
    )
    assert classify_doc_kind(bl) == "bl"


def test_normal_si_is_not_structurally_wrong():
    fields = _full_fields()
    si = ExtractedDocument(
        filename="si.txt",
        text="SHIPPING INSTRUCTION\nShipper ACME\nConsignee BETA\n" + "x" * 80,
        fields=fields,
    )
    assert classify_doc_kind(si) == "si"


def test_ocr_failed_bl_is_not_wrong_doc_type():
    noise = ("scan noise ## " * 40).strip()
    bl = ExtractedDocument(filename="scan.pdf", text=noise, fields={})
    assert classify_doc_kind(bl) == "unknown"
    email = EmailMessage(email_id="e1", subject="SI vs BL", attachment_paths=["si.txt", "scan.pdf"])
    si = ExtractedDocument(
        filename="si.txt",
        kind="si",
        text="SHIPPING INSTRUCTION\n" + "x" * 80,
        fields=_full_fields(),
    )
    result = health_check(email, [si, bl], si=si, bl=bl)
    assert result.reason != ReviewReason.WRONG_DOC_TYPE


def test_garbage_content_does_not_crash():
    junk = ExtractedDocument(filename="a.txt", text=("xyz " * 40).strip(), fields={})
    junk2 = ExtractedDocument(filename="b.txt", text=("qwe " * 40).strip(), fields={})
    email = EmailMessage(email_id="e1", subject="SI vs BL", attachment_paths=["a.txt", "b.txt"])
    result = health_check(email, [junk, junk2])
    assert result.reason != ReviewReason.WRONG_DOC_TYPE
    assert result.reason in {
        None,
        ReviewReason.MISSING_VALUE,
        ReviewReason.UNREADABLE,
        ReviewReason.MISSING_ATTACHMENT,
    }


def test_bl_filename_on_commercial_invoice_is_wrong_doc_type():
    invoice_text = (
        "COMMERCIAL INVOICE\nInvoice No.: 5250078266\n"
        "Description  Qty  Unit Price  Amount (USD)\n"
        "Paper  500  45.00  22500\n"
        "*** THIS IS A COMMERCIAL INVOICE - NOT A SHIPPING INSTRUCTION ***\n"
    )
    inv = ExtractedDocument(filename="email_pair_BL.txt", text=invoice_text, fields={})
    assert classify_doc_kind(inv) == "other"
    fields = _full_fields()
    si = ExtractedDocument(
        filename="email_pair_SI.txt",
        kind="si",
        text="SHIPPING INSTRUCTION\n" + "x" * 80,
        fields=fields,
    )
    inv.kind = classify_doc_kind(inv)
    email = EmailMessage(
        email_id="wd1",
        subject="TO CONFIRM DOCS",
        body_text="Please find attached the SI and the Commercial Invoice. Kindly confirm the BL is in order.",
        attachment_paths=["email_pair_SI.txt", "email_pair_BL.txt"],
    )
    result = health_check(email, [si, inv], si=si, bl=inv)
    assert result.ok is False
    assert result.reason == ReviewReason.WRONG_DOC_TYPE


def test_bl_filename_on_packing_list_is_wrong_doc_type():
    pl_text = (
        "PACKING LIST\nShipper: ACME\nConsignee: BETA\n"
        "CTN-001  10kg\n*** PACKING LIST ONLY - NO PORT OR VESSEL DETAILS ***\n"
    )
    pl = ExtractedDocument(
        filename="email_pair_BL.txt",
        text=pl_text,
        fields={"shipper": _fv("shipper", "ACME"), "consignee": _fv("consignee", "BETA")},
    )
    assert classify_doc_kind(pl) == "other"


def test_real_bl_filename_stays_bl_even_if_invoice_mentioned():
    fields = _full_fields()
    bl = ExtractedDocument(
        filename="email_pair_BL.txt",
        text="BILL OF LADING\nFreight payable as per commercial invoice.\nPort of loading Shanghai\n",
        fields=fields,
    )
    assert classify_doc_kind(bl) == "bl"


def test_si_listing_packing_list_as_required_doc_stays_si():
    fields = _full_fields()
    si = ExtractedDocument(
        filename="email_pair_SI.txt",
        text="SHIPPING INSTRUCTION\nDocuments Required:\n1) 3 Original invoice\n2) 3 Packing list\n",
        fields=fields,
    )
    assert classify_doc_kind(si) == "si"


def test_priority_when_multiple_conditions_hit():
    fields = _full_fields()

    # Attachment missing beats a damaged file: if the pair never arrived,
    # there is nothing whose OCR quality or emptiness can be judged.
    missing_email = EmailMessage(
        email_id="p1",
        subject="please compare the draft BL",
        body_text="Please compare the draft BL. The attachment is missing and was not received.",
        attachment_paths=["si.pdf"],
    )
    empty_si = ExtractedDocument(filename="si.pdf", kind="si", text="", parser_used="vision")
    r1 = health_check(missing_email, [empty_si])
    assert r1.reason == ReviewReason.MISSING_ATTACHMENT

    # Unreadable beats wrong type: a scan with no text layer cannot be
    # classified as an invoice even if a fragment happens to say so.
    si = ExtractedDocument(
        filename="si.txt",
        kind="si",
        text="SHIPPING INSTRUCTION\n" + "x" * 80,
        fields=fields,
    )
    faint_invoice = ExtractedDocument(
        filename="bl.pdf",
        kind="other",
        text="commercial invoice",
        parser_used="vision",
        fields={},
    )
    r2 = health_check(
        EmailMessage(email_id="p2", subject="SI vs BL", attachment_paths=["si.txt", "bl.pdf"]),
        [si, faint_invoice],
        si=si,
        bl=faint_invoice,
    )
    assert r2.reason == ReviewReason.UNREADABLE

    # Wrong type beats a missing field: if the file is not an SI/BL,
    # asking which compare field is blank is meaningless.
    si_tba = ExtractedDocument(
        filename="si.txt",
        kind="si",
        text="SHIPPING INSTRUCTION\n" + "x" * 80,
        fields={**fields, "shipper": _fv("shipper", "TBA")},
    )
    invoice = ExtractedDocument(
        filename="inv.pdf",
        kind="other",
        text="COMMERCIAL INVOICE\n" + "x" * 80,
        fields=fields,
    )
    r3 = health_check(
        EmailMessage(email_id="p3", subject="SI vs BL", attachment_paths=["si.txt", "inv.pdf"]),
        [si_tba, invoice],
        si=si_tba,
        bl=invoice,
    )
    assert r3.reason == ReviewReason.WRONG_DOC_TYPE


def test_two_line_to_the_order_of_is_not_missing_consignee():
    from harbormaster.reader.extractor import FieldExtractor

    text = (
        "SHIPPING INSTRUCTION\n"
        "Shipper (Principal or Seller)\n"
        "ACME TRADING\n"
        "To the Order of\n"
        "BETA LOGISTICS\n"
        "Notify\n"
        "BETA LOGISTICS\n"
        "PORT OF LOADING\n"
        "SINGAPORE\n"
        "POD\n"
        "ROTTERDAM\n"
        "Total Containers\n"
        "3 x 40HC\n"
        "GROSS WEIGHT\n"
        "1000\n"
    )
    si = FieldExtractor(degrade=True, rules_only=True).extract_text(text, filename="si.txt")
    assert "consignee" in si.fields
    assert "BETA" in si.fields["consignee"].raw_value.upper()


def test_docx_table_fields_are_not_missing_value(tmp_path):
    from docx import Document

    from harbormaster.reader.extractor import FieldExtractor

    path = tmp_path / "email_pair_BL.docx"
    doc = Document()
    doc.add_paragraph("BILL OF LADING (DRAFT)")
    table = doc.add_table(rows=7, cols=2)
    rows = [
        ("Shipper (Principal or Seller)", "ACME TRADING"),
        ("Consignee", "BETA LOGISTICS"),
        ("Notify", "BETA LOGISTICS"),
        ("PORT OF LOADING", "SINGAPORE"),
        ("POD", "ROTTERDAM"),
        ("Total Containers", "3 x 40HC"),
        ("Gross Wt (kgs)", "1000"),
    ]
    for i, (label, value) in enumerate(rows):
        table.rows[i].cells[0].text = label
        table.rows[i].cells[1].text = value
    doc.save(path)

    extracted = FieldExtractor(degrade=True, rules_only=True).extract_path(path)
    assert extracted.kind == "bl"
    for name in (
        "shipper",
        "consignee",
        "notify_party",
        "port_of_loading",
        "port_of_discharge",
        "container_count",
        "gross_weight_kg",
    ):
        assert name in extracted.fields, extracted.fields.keys()

    si = ExtractedDocument(
        filename="email_pair_SI.txt",
        kind="si",
        text="SHIPPING INSTRUCTION\n" + "x" * 80,
        fields=_full_fields(),
    )
    email = EmailMessage(
        email_id="docx1",
        subject="SI vs BL",
        attachment_paths=["email_pair_SI.txt", str(path)],
    )
    result = health_check(email, [si, extracted], si=si, bl=extracted)
    assert result.reason != ReviewReason.MISSING_VALUE
    assert result.ok is True or result.reason is None


def test_parenthetical_pol_placeholder_is_missing_value():
    """'Port of Loading (POL): ____MT' must extract the blank, not '(POL): ____MT'."""
    from harbormaster.reader.extractor import FieldExtractor

    si_text = (
        "SHIPPING INSTRUCTION\n"
        "Shipper: ACME\n"
        "Consignee: BETA\n"
        "Notify: BETA\n"
        "Port of Loading (POL): ____MT\n"
        "Port of Discharge (POD): TBA\n"
        "Total Containers: 2 x 20'GP\n"
        "GROSS WEIGHT: 1000 KG\n"
    )
    bl_text = (
        "BILL OF LADING\n"
        "Shipper: ACME\n"
        "Consignee: BETA\n"
        "Notify: BETA\n"
        "Port of Loading: SINGAPORE\n"
        "Port of Discharge: ROTTERDAM\n"
        "No. of Containers: 2 x 20'GP\n"
        "GROSS WEIGHT: 1000 KG\n"
    )
    extractor = FieldExtractor(degrade=True, rules_only=True)
    si = extractor.extract_text(si_text, filename="si.txt")
    bl = extractor.extract_text(bl_text, filename="bl.txt")
    assert si.fields["port_of_loading"].raw_value == "____MT"
    assert si.fields["port_of_discharge"].raw_value == "TBA"
    email = EmailMessage(
        email_id="blank_pol",
        subject="Please compare the SI and draft BL",
        attachment_paths=["si.txt", "bl.txt"],
    )
    result = health_check(email, [si, bl], si=si, bl=bl)
    assert result.ok is False
    assert result.reason == ReviewReason.MISSING_VALUE
    assert "port_of_loading" in result.missing_fields
    assert "port_of_discharge" in result.missing_fields


def test_chinese_gross_weight_placeholder_is_missing_value():
    from harbormaster.reader.extractor import FieldExtractor

    si_text = (
        "SHIPPING INSTRUCTION\n"
        "Shipper: ACME\n"
        "Consignee: BETA\n"
        "Notify: BETA\n"
        "PORT OF LOADING: SINGAPORE\n"
        "POD: ROTTERDAM\n"
        "No. of Containers or Packages: 1 x 20'GP\n"
        "Gross Weight毛重(KGS): N/A\n"
    )
    bl_text = (
        "BILL OF LADING\n"
        "Shipper: ACME\n"
        "Consignee: BETA\n"
        "Notify: BETA\n"
        "Port of Loading: SINGAPORE\n"
        "Port of Discharge: ROTTERDAM\n"
        "No. of Containers: 1 x 20'GP\n"
        "GROSS WEIGHT: 900 KG\n"
    )
    extractor = FieldExtractor(degrade=True, rules_only=True)
    si = extractor.extract_text(si_text, filename="si.txt")
    bl = extractor.extract_text(bl_text, filename="bl.txt")
    assert si.fields["gross_weight_kg"].raw_value == "N/A"
    email = EmailMessage(
        email_id="blank_wt",
        subject="Please compare the SI and draft BL",
        attachment_paths=["si.txt", "bl.txt"],
    )
    result = health_check(email, [si, bl], si=si, bl=bl)
    assert result.ok is False
    assert result.reason == ReviewReason.MISSING_VALUE
    assert "gross_weight_kg" in result.missing_fields


